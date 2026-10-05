from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file
import sqlite3, os, io
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import load_workbook, Workbook

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "change-this-secret-key")
DB = os.environ.get("DB_PATH", "payments.db")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "change-me-before-deploy")

def db():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    return con

def init_db():
    con = db()
    con.execute("""CREATE TABLE IF NOT EXISTS payments(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        amount REAL NOT NULL,
        date TEXT DEFAULT '',
        note TEXT DEFAULT ''
    )""")
    con.commit()
    con.close()

def logged_in():
    return session.get("admin") is True

@app.route("/")
def index():
    con = db()
    rows = con.execute("SELECT * FROM payments ORDER BY id DESC").fetchall()
    total = con.execute("SELECT COALESCE(SUM(amount),0) FROM payments").fetchone()[0]
    con.close()
    return render_template("index.html", rows=rows, total=total, admin=logged_in())

@app.route("/login", methods=["GET","POST"])
def login():
    if request.method == "POST":
        password = request.form.get("password","")
        if check_password_hash(generate_password_hash(ADMIN_PASSWORD), password):
            session["admin"] = True
            return redirect(url_for("index"))
        flash("गलत password")
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("index"))

@app.route("/add", methods=["POST"])
def add():
    if not logged_in(): return redirect(url_for("login"))
    name = request.form.get("name","").strip()
    amount = request.form.get("amount","0")
    date = request.form.get("date","")
    note = request.form.get("note","")
    try:
        amount = float(amount)
    except ValueError:
        amount = 0
    if name:
        con = db()
        con.execute("INSERT INTO payments(name,amount,date,note) VALUES(?,?,?,?)",
                    (name, amount, date, note))
        con.commit(); con.close()
    return redirect(url_for("index"))

@app.route("/delete/<int:id>", methods=["POST"])
def delete(id):
    if not logged_in(): return redirect(url_for("login"))
    con = db(); con.execute("DELETE FROM payments WHERE id=?", (id,))
    con.commit(); con.close()
    return redirect(url_for("index"))

@app.route("/upload", methods=["POST"])
def upload():
    if not logged_in(): return redirect(url_for("login"))
    f = request.files.get("excel")
    if not f or not f.filename:
        flash("Excel file चुनिए")
        return redirect(url_for("index"))
    try:
        wb = load_workbook(f, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            flash("Excel खाली है")
            return redirect(url_for("index"))
        headers = [str(x).strip().lower() if x is not None else "" for x in rows[0]]
        def col(*names):
            for n in names:
                if n in headers: return headers.index(n)
            return None
        ni, ai, di, oi = col("name","नाम"), col("amount","रकम","amount (rs)","rupees"), col("date","तारीख"), col("note","नोट")
        if ni is None or ai is None:
            flash("पहली row में Name और Amount columns होना जरूरी है")
            return redirect(url_for("index"))
        con = db()
        for r in rows[1:]:
            if ni >= len(r) or ai >= len(r) or r[ni] is None or r[ai] is None: continue
            try: amount = float(r[ai])
            except: continue
            date = str(r[di]) if di is not None and di < len(r) and r[di] is not None else ""
            note = str(r[oi]) if oi is not None and oi < len(r) and r[oi] is not None else ""
            con.execute("INSERT INTO payments(name,amount,date,note) VALUES(?,?,?,?)",
                        (str(r[ni]), amount, date, note))
        con.commit(); con.close()
        flash("Excel data जोड़ दिया गया")
    except Exception as e:
        flash("Excel पढ़ने में समस्या हुई: " + str(e))
    return redirect(url_for("index"))

@app.route("/export")
def export():
    if not logged_in(): return redirect(url_for("login"))
    con = db(); rows = con.execute("SELECT name,amount,date,note FROM payments ORDER BY id").fetchall(); con.close()
    wb = Workbook(); ws = wb.active
    ws.append(["Name","Amount","Date","Note"])
    for r in rows: ws.append([r["name"], r["amount"], r["date"], r["note"]])
    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio, as_attachment=True, download_name="payments_updated.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT",5000)))
